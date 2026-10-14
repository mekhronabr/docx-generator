"""Fill the supplied Talabot form without rebuilding its layout or ZIP parts."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import re
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock

TEMPLATE = Path(__file__).with_name('template_talabot.docx')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def lines(text):
    return [line.strip() for line in text.splitlines() if line.strip()]


def display_date(value):
    text = str(value).strip()
    return pd.to_datetime(text, dayfirst=not bool(re.match(r'^\d{4}-\d{2}-\d{2}', text)), errors='raise').strftime('%d.%m.%Y')


def apply_style(run, style):
    props = run.find(W + 'rPr')
    if props is None:
        props = etree.Element(W + 'rPr')
        run.insert(0, props)
    for key, tag in [('bold', 'b'), ('italic', 'i'), ('strike', 'strike')]:
        if key in style:
            node = props.find(W + tag)
            if node is None:
                node = etree.SubElement(props, W + tag)
            node.set(W + 'val', '1' if style[key] else '0')
    if 'underline' in style:
        node = props.find(W + 'u')
        if node is None:
            node = etree.SubElement(props, W + 'u')
        node.set(W + 'val', style['underline'] or 'none')
    if style.get('color'):
        node = props.find(W + 'color')
        if node is None:
            node = etree.SubElement(props, W + 'color')
        node.set(W + 'val', style['color'])
    if style.get('font'):
        node = props.find(W + 'rFonts')
        if node is None:
            node = etree.SubElement(props, W + 'rFonts')
        for attr in ('ascii', 'hAnsi', 'cs'):
            node.set(W + attr, style['font'])
    if style.get('vertical'):
        etree.SubElement(props, W + 'vertAlign').set(W + 'val', style['vertical'])
    if style.get('size'):
        for tag in ('sz', 'szCs'):
            node = props.find(W + tag)
            if node is None:
                node = etree.SubElement(props, W + tag)
            node.set(W + 'val', str(round(style['size'] * 2)))


def fill_slot(control, text, fragments=None):
    content = control.find(W + 'sdtContent')
    base_run = content.find(W + 'r')
    if base_run is None:
        raise ValueError('Неподдерживаемое поле шаблона.')
    parent = control.getparent()
    position = parent.index(control)
    # Rich fragments are valid only while the user has kept the original text.
    if not fragments or ''.join(f['text'] for f in fragments) != text:
        style = fragments[0].get('style', {}) if fragments else {}
        fragments = [{'text': text, 'style': style}]
    for fragment in fragments:
        run = deepcopy(base_run)
        for node in list(run):
            if node.tag != W + 'rPr':
                run.remove(node)
        apply_style(run, fragment.get('style', {}))
        for index, line in enumerate(fragment['text'].split('\n')):
            if index:
                etree.SubElement(run, W + 'br')
            node = etree.SubElement(run, W + 't')
            node.text = line
            node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        parent.insert(position, run)
        position += 1
    parent.remove(control)


def render_root(data, template_bytes):
    students, absent = data['students'], data.get('absent', [])
    with ZipFile(BytesIO(template_bytes)) as source:
        root = etree.fromstring(source.read('word/document.xml'))
    absence_tags = [control.find(W + 'sdtPr/' + W + 'tag').get(W + 'val') for control in root.iter(W + 'sdt') if control.find(W + 'sdtPr/' + W + 'tag').get(W + 'val').startswith('absent_')]
    if not 1 <= len(students) <= 40:
        raise ValueError('В основном списке должно быть от 1 до 40 студентов.')
    if len(absent) > len(absence_tags):
        raise ValueError(f'В списке отсутствующих допустимо до {len(absence_tags)} студентов — по числу мест в шаблоне.')
    for field in ('form_code', 'group', 'date', 'curator', 'deputy'):
        if not data.get(field, '').strip():
            raise ValueError('Заполните код факультета, группу, дату и ФИО ответственных.')
    slots = {key: data[key] for key in ('form_code', 'group', 'date', 'curator', 'deputy')}
    slots.update({f'student_{i + 1}': students[i] if i < len(students) else '' for i in range(40)})
    slots.update({tag: absent[i] if i < len(absent) else '' for i, tag in enumerate(absence_tags)})
    if any(re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value) for value in slots.values()):
        raise ValueError('В тексте обнаружены недопустимые управляющие символы.')
    formats = data.get('formats', {})
    formats = dict(formats)
    for i, tag in reversed(list(enumerate(absence_tags))):
        formats[tag] = data.get('formats', {}).get(f'absent_{i + 1}', [])
    for control in list(root.iter(W + 'sdt')):
        tag = control.find(W + 'sdtPr/' + W + 'tag').get(W + 'val')
        fill_slot(control, slots[tag], formats.get(tag))
    compact_document_tail(root)
    return root


def compact_document_tail(root):
    """Keep Word's final table paragraph from overflowing onto an empty page."""
    for paragraph in root.iter(W + 'p'):
        if ''.join(paragraph.itertext()).strip() or any(paragraph.find('.//' + W + tag) is not None for tag in ('br', 'drawing', 'pict', 'pageBreakBefore')):
            continue
        props = paragraph.find(W + 'pPr')
        if props is None:
            props = etree.Element(W + 'pPr'); paragraph.insert(0, props)
        spacing = props.find(W + 'spacing')
        if spacing is None: spacing = etree.SubElement(props, W + 'spacing')
        for key, value in [('before', '0'), ('after', '0'), ('line', '20'), ('lineRule', 'exact')]:
            spacing.set(W + key, value)
        font = props.find(W + 'rPr')
        if font is None: font = etree.SubElement(props, W + 'rPr')
        size = font.find(W + 'sz')
        if size is None: size = etree.SubElement(font, W + 'sz')
        size.set(W + 'val', '2')
    body = root.find(W + 'body')
    section = body.find(W + 'sectPr')
    position = body.index(section) if section is not None else len(body)
    while position and body[position - 1].tag == W + 'p':
        paragraph = body[position - 1]
        if ''.join(paragraph.itertext()).strip() or paragraph.find('.//' + W + 'br') is not None:
            break
        body.remove(paragraph)
        position -= 1
    paragraph = etree.Element(W + 'p')
    props = etree.SubElement(paragraph, W + 'pPr')
    spacing = etree.SubElement(props, W + 'spacing')
    for key, value in [('before', '0'), ('after', '0'), ('line', '20'), ('lineRule', 'exact')]:
        spacing.set(W + key, value)
    font = etree.SubElement(props, W + 'rPr')
    etree.SubElement(font, W + 'sz').set(W + 'val', '2')
    body.insert(position, paragraph)


def write_package(root, template_bytes):
    output = BytesIO()
    with ZipFile(BytesIO(template_bytes)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            contents = etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True) if item.filename == 'word/document.xml' else source.read(item.filename)
            target.writestr(item, contents)
    return output.getvalue()


def generate(data, template_bytes=None):
    template_bytes = template_bytes or TEMPLATE.read_bytes()
    return write_package(render_root(data, template_bytes), template_bytes)


def main():
    st.write('Лист ознакомления группы с требованиями. ФИО вводятся по одному на строку или загружаются из Excel.')
    if not TEMPLATE.exists():
        st.error('Отсутствует template_talabot.docx рядом с app.py.')
        return
    version = hashlib.sha256(TEMPLATE.read_bytes() + Path(__file__).read_bytes()).hexdigest()
    if st.session_state.get('talabot_version') != version:
        st.session_state.pop('talabot_result', None)
        st.session_state.pop('talabot_batch_result', None)
        st.session_state.talabot_version = version
    excel_example = BytesIO()
    pd.DataFrame(columns=['Ному насаби донишҷӯ', 'Гурӯҳ', 'Код факультета', 'Дата', 'ФИО Куратора', 'ФИО Заместителя', 'Отсутствует']).to_excel(excel_example, index=False)
    st.download_button('Скачать пустую таблицу Excel для Талабот', excel_example.getvalue(), file_name='Талабот_ҷадвали_холӣ.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', on_click='ignore')
    mode = st.radio('Режим генерации', ['Одна группа', 'Несколько групп'], horizontal=True)
    from examples_ui import show_examples
    show_examples('talabot_batch' if mode == 'Несколько групп' else 'talabot', TEMPLATE)
    if mode == 'Несколько групп':
        batch_main()
        return
    upload = st.file_uploader('Excel со списком студентов (.xlsx)', type=['xlsx'], key='talabot_excel')
    imported = ''
    imported_formats = {}
    metadata_formats = {}
    defaults = dict(form_code='ФЭ_11.', group='1_53.01.05тб', curator='', deputy='', date='10.10.2025', absent='')
    if upload is not None:
        try:
            if upload.size > 20 * 1024 * 1024:
                raise ValueError('Размер Excel должен быть не более 20 МБ.')
            with ZipFile(BytesIO(upload.getvalue())) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('Распакованный Excel превышает 100 МБ.')
            df = pd.read_excel(BytesIO(upload.getvalue()), dtype=str, keep_default_na=False, nrows=100001)
            if len(df) > 100000:
                raise ValueError('Допустимо до 100 000 строк. Разделите Excel.')
            if not len(df.columns):
                raise ValueError('В Excel нет колонок.')
            df.attrs['formats'] = excel_formats(upload.getvalue(), df)
            columns = list(df.columns)
            preferred = next((i for i, c in enumerate(columns) if str(c).strip().casefold() in ('фио студента', 'ному насаби донишҷӯ', 'фио', 'ному насаб')), 0)
            column = st.selectbox('Колонка с ФИО студентов', columns, index=preferred, key='talabot_name_column')
            group_columns = [c for c in df.columns if str(c).strip().casefold() in ('группа', 'гурӯҳ', 'гурух')]
            if group_columns:
                group_column = group_columns[0]
                selected = st.selectbox('Группа из Excel', list(dict.fromkeys(df[group_column])), key='talabot_excel_group')
                df = df.loc[df[group_column] == selected]
            imported = '\n'.join(str(v).strip() for v in df[column] if str(v).strip())
            imported_formats = {str(df.at[i, column]).strip(): df.attrs['formats'].get((i, column), []) for i in df.index if str(df.at[i, column]).strip()}
            for tag, candidates in [('form_code', ['Код факультета']), ('group', ['Группа', 'Гурӯҳ']), ('date', ['Дата']), ('curator', ['ФИО Куратора', 'Куратори гурӯҳ']), ('deputy', ['ФИО Заместителя', 'ФИО Зам. декана', 'Муовини декан оид ба тарбия'])]:
                field = next((c for c in candidates if c in df), None)
                if field is not None:
                    indices = [i for i in df.index if str(df.at[i, field]).strip()]
                    if indices:
                        defaults[tag] = str(df.at[indices[0], field]).strip()
                        metadata_formats[tag] = df.attrs['formats'].get((indices[0], field), [])
            defaults['date'] = display_date(defaults['date'])
            if 'Отсутствует' in df:
                defaults['absent'] = '\n'.join(str(df.at[i, column]).strip() for i in df.index if str(df.at[i, 'Отсутствует']).strip().casefold() in {'да', '1', 'yes', 'true', '+'})
        except Exception as error:
            st.error(f'Не удалось прочитать Excel: {error}')
            return
    with st.form('talabot_form'):
        code = st.text_input('Код факультета', value=defaults['form_code'])
        group = st.text_input('Группа', value=defaults['group'])
        when = st.date_input('Дата ознакомления', value=pd.to_datetime(defaults['date'], format='%d.%m.%Y').date())
        curator = st.text_input('Фамилия и имя — куратори гурӯҳи академӣ', value=defaults['curator'])
        deputy = st.text_input('Фамилия и имя — муовини директор оид ба тарбия', value=defaults['deputy'])
        students = st.text_area('Фамилии и имена студентов (до 40)', value=imported, height=300)
        absent = st.text_area('Отсутствующие (до 10, по одному на строку)', value=defaults['absent'], height=120)
        with st.expander('Оформление полей при ручном вводе'):
            st.caption('Из Excel сохраняются жирность, курсив, подчёркивание и регистр. Эти настройки применяются к тексту, введённому вручную.')
            manual_styles = {}
            for tag, label, defaults in [('form_code', 'Код факультета', (True, False, False)), ('group', 'Группа', (True, True, True)), ('date', 'Дата', (False, True, True)), ('students', 'ФИО студентов', (False, False, False)), ('absent', 'Отсутствующие', (False, False, False)), ('curator', 'Куратор', (True, True, True)), ('deputy', 'Заместитель', (True, True, True))]:
                st.write(label)
                controls = st.columns(3)
                manual_styles[tag] = dict(bold=controls[0].checkbox('Жирный', value=defaults[0], key=f'talabot_{tag}_bold'), italic=controls[1].checkbox('Курсив', value=defaults[1], key=f'talabot_{tag}_italic'), underline='single' if controls[2].checkbox('Подчёркнутый', value=defaults[2], key=f'talabot_{tag}_underline') else 'none')
        submitted = st.form_submit_button('Сгенерировать документ', type='primary')
    if submitted:
        st.session_state.pop('talabot_result', None)
        try:
            data = dict(form_code=code.strip(), group=group.strip(), date=when.strftime('%d.%m.%Y'), curator=curator.strip(), deputy=deputy.strip(), students=lines(students), absent=lines(absent))
            data['formats'] = {tag: metadata_formats.get(tag) or [{'text': data[tag], 'style': manual_styles[tag]}] for tag in ('form_code', 'group', 'date', 'curator', 'deputy')}
            for field, prefix in [('students', 'student'), ('absent', 'absent')]:
                for i, name in enumerate(data[field], 1):
                    data['formats'][f'{prefix}_{i}'] = imported_formats.get(name) or [{'text': name, 'style': manual_styles[field]}]
            st.session_state.talabot_result = generate(data)
        except Exception as error:
            st.error(f'Не удалось создать документ: {error}')
    if st.session_state.get('talabot_result'):
        st.download_button('Скачать заполненный Талабот', st.session_state.talabot_result, file_name='Талабот.docx', mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document', on_click='ignore')


BATCH_FIELDS = ['Группа', 'Код факультета', 'Дата', 'ФИО Куратора', 'ФИО Заместителя']


def excel_formats(contents, df):
    """Read font emphasis and rich-text fragments that pandas intentionally drops."""
    workbook = load_workbook(BytesIO(contents), data_only=True, rich_text=True)
    sheet = workbook.worksheets[0]
    default_font = workbook._fonts[0]
    formats = {}

    def font_style(font, baseline):
        result = {}
        for attr, key in [('b', 'bold'), ('i', 'italic'), ('strike', 'strike')]:
            if getattr(font, attr, None) != getattr(baseline, attr, None):
                result[key] = bool(getattr(font, attr, None))
        underline = getattr(font, 'u', None)
        if underline != getattr(baseline, 'u', None):
            result['underline'] = {'singleAccounting': 'single', 'doubleAccounting': 'double'}.get(underline, underline) or 'none'
        color = getattr(font, 'color', None)
        if color is not None and color.type == 'rgb':
            result['color'] = color.rgb[-6:]
        font_name = next((value for value in (getattr(font, 'name', None), getattr(font, 'rFont', None)) if isinstance(value, str)), None)
        baseline_name = next((value for value in (getattr(baseline, 'name', None), getattr(baseline, 'rFont', None)) if isinstance(value, str)), None)
        if font_name and font_name != baseline_name:
            result['font'] = font_name
        if getattr(font, 'vertAlign', None):
            result['vertical'] = font.vertAlign
        size = getattr(font, 'sz', None)
        if isinstance(size, (int, float)) and size != getattr(baseline, 'sz', None):
            result['size'] = size
        return result

    for index in df.index:
        for col_index, column in enumerate(df.columns, 1):
            cell = sheet.cell(int(index) + 2, col_index)
            base_style = font_style(cell.font, default_font)
            if str(column).casefold() in ('фио студента', 'ному насаби донишҷӯ', 'фио', 'ному насаб'):
                base_style.update(bold=bool(cell.font.b), italic=bool(cell.font.i), underline=cell.font.u or 'none')
            if isinstance(cell.value, CellRichText):
                fragments = []
                for part in cell.value:
                    if isinstance(part, TextBlock):
                        style = {**base_style, **font_style(part.font, cell.font)}
                        fragments.append({'text': part.text, 'style': style})
                    else:
                        fragments.append({'text': str(part), 'style': base_style})
            else:
                fragments = [{'text': str(df.at[index, column]), 'style': base_style}]
            formats[(index, column)] = fragments
    workbook.close()
    return formats


def read_batch_excel(contents):
    if len(contents) > 20 * 1024 * 1024:
        raise ValueError('Размер Excel должен быть не более 20 МБ.')
    with ZipFile(BytesIO(contents)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Распакованный Excel превышает 100 МБ.')
    df = pd.read_excel(BytesIO(contents), dtype=str, keep_default_na=False, nrows=100001)
    if len(df) > 100000:
        raise ValueError('Допустимо до 100 000 строк. Разделите Excel.')
    aliases = {
        'ному насаби донишҷӯ': 'ФИО Студента', 'фио': 'ФИО Студента',
        'ному насаб': 'ФИО Студента', 'гурӯҳ': 'Группа', 'гурух': 'Группа',
        'куратори гурӯҳ': 'ФИО Куратора',
        'фио зам. декана': 'ФИО Заместителя',
        'муовини декан оид ба тарбия': 'ФИО Заместителя',
        'муовини директор оид ба тарбия': 'ФИО Заместителя',
    }
    aliases.update({c.casefold(): c for c in ['ФИО Студента', *BATCH_FIELDS, 'Отсутствует']})
    df.columns = [aliases.get(str(c).strip().casefold(), str(c).strip()) for c in df.columns]
    if df.columns.duplicated().any():
        raise ValueError('В Excel повторяются названия колонок.')
    for column in ('ФИО Студента', 'Группа'):
        if column not in df:
            raise ValueError(f'В Excel отсутствует колонка «{column}».')
    df = df.map(lambda value: str(value).strip())
    df = df.loc[df.ne('').any(axis=1)].copy()
    if df.empty:
        raise ValueError('В Excel нет студентов.')
    bad = df[['ФИО Студента', 'Группа']].eq('').any(axis=1)
    if bad.any():
        raise ValueError('Заполните ФИО и группу в строках: ' + ', '.join(str(i + 2) for i in df.index[bad][:20]))
    if 'Отсутствует' in df:
        allowed = {'', 'да', 'нет', '1', '0', 'yes', 'no', 'true', 'false', '+', '-'}
        if not df['Отсутствует'].str.casefold().isin(allowed).all():
            raise ValueError('В колонке «Отсутствует» укажите «да», «нет» или оставьте пусто.')
    df.attrs['formats'] = excel_formats(contents, df)
    return df


def batch_settings(df):
    settings = []
    for group, rows in df.groupby('Группа', sort=False):
        item = {'Группа': group}
        for field in BATCH_FIELDS[1:]:
            values = list(dict.fromkeys(value for value in rows[field] if value)) if field in rows else []
            if len(values) > 1:
                raise ValueError(f'Группа {group}: разные значения в колонке «{field}». Укажите одно значение для группы.')
            value = values[0] if values else ''
            if field == 'Дата' and value:
                value = display_date(value)
            item[field] = value
        settings.append(item)
    return pd.DataFrame(settings, columns=BATCH_FIELDS)


def generate_batch(df, settings, progress=lambda done, total: None):
    """One shared Word document; every group starts on a fresh page."""
    if len(settings) != len(df['Группа'].unique()) or settings['Группа'].duplicated().any() or set(settings['Группа']) != set(df['Группа']):
        raise ValueError('Настройки должны содержать каждую группу ровно один раз.')
    template_bytes = TEMPLATE.read_bytes()
    with ZipFile(BytesIO(template_bytes)) as source:
        combined = etree.fromstring(source.read('word/document.xml'))
    body = combined.find(W + 'body')
    section = deepcopy(body.find(W + 'sectPr'))
    for child in list(body):
        body.remove(child)
    formats = df.attrs.get('formats', {})
    total_size = 0
    for index, item in enumerate(settings.to_dict('records'), 1):
        group = item['Группа']
        rows = df.loc[df['Группа'] == group]
        missing = [field for field in BATCH_FIELDS if pd.isna(item[field]) or not str(item[field]).strip()]
        if missing:
            raise ValueError(f'Группа {group}: заполните ' + ', '.join(missing))
        when = display_date(item['Дата'])
        absent = rows.loc[rows['Отсутствует'].str.casefold().isin({'да', '1', 'yes', 'true', '+'}), 'ФИО Студента'].tolist() if 'Отсутствует' in rows else []
        data = dict(form_code=str(item['Код факультета']).strip(), group=group, date=when, curator=str(item['ФИО Куратора']).strip(), deputy=str(item['ФИО Заместителя']).strip(), students=rows['ФИО Студента'].tolist(), absent=absent)
        data['formats'] = {}
        for field, slot in [('Код факультета', 'form_code'), ('Группа', 'group'), ('Дата', 'date'), ('ФИО Куратора', 'curator'), ('ФИО Заместителя', 'deputy')]:
            source_index = next((i for i in rows.index if field in rows and rows.at[i, field]), rows.index[0])
            data['formats'][slot] = formats.get((source_index, field), [])
        for position, source_index in enumerate(rows.index, 1):
            data['formats'][f'student_{position}'] = formats.get((source_index, 'ФИО Студента'), [])
        if 'Отсутствует' in rows:
            absent_indices = rows.index[rows['Отсутствует'].str.casefold().isin({'да', '1', 'yes', 'true', '+'})]
            for position, source_index in enumerate(absent_indices, 1):
                data['formats'][f'absent_{position}'] = formats.get((source_index, 'ФИО Студента'), [])
        try:
            page_root = render_root(data, template_bytes)
        except ValueError as error:
            raise ValueError(f'Группа {group}: {error}') from error
        if index > 1:
            paragraph = etree.SubElement(body, W + 'p')
            props = etree.SubElement(paragraph, W + 'pPr')
            spacing = etree.SubElement(props, W + 'spacing')
            spacing.set(W + 'before', '0')
            spacing.set(W + 'after', '0')
            spacing.set(W + 'line', '20')
            spacing.set(W + 'lineRule', 'exact')
            etree.SubElement(props, W + 'pageBreakBefore')
        page_body = page_root.find(W + 'body')
        for child in list(page_body):
            if child.tag != W + 'sectPr':
                body.append(child)
        total_size += len(etree.tostring(page_root))
        if total_size > 64 * 1024 * 1024:
            raise ValueError('Документ слишком большой. Разделите Excel на несколько файлов.')
        progress(index, len(settings))
    body.append(section)
    return write_package(combined, template_bytes)


def batch_main():
    st.write('Загрузите Excel со всеми группами. Одна строка — один студент; колонка группы обязательна.')
    st.caption('Данные ответственных можно заполнить в Excel или в настройках ниже. «Отсутствует»: да/нет; отсутствующий остаётся и в основном списке.')
    upload = st.file_uploader('Excel для нескольких групп (.xlsx)', type=['xlsx'], key='talabot_batch_excel')
    if upload is None:
        return
    try:
        df = read_batch_excel(upload.getvalue())
        settings = batch_settings(df)
    except Exception as error:
        st.error(f'Не удалось прочитать Excel: {error}')
        return
    selected = st.multiselect('Группы для генерации', settings['Группа'].tolist(), default=settings['Группа'].tolist())
    if not selected:
        st.info('Выберите хотя бы одну группу.')
        return
    df = df.loc[df['Группа'].isin(selected)]
    settings = settings.loc[settings['Группа'].isin(selected)]
    st.info(f'Студентов: {len(df)}. Групп: {len(settings)}. Результат — один Word.')
    with st.form('talabot_batch_form'):
        st.write('Настройки каждого документа')
        edited = st.data_editor(settings, hide_index=True, disabled=['Группа'], num_rows='fixed', key='talabot_batch_settings')
        submitted = st.form_submit_button('Сгенерировать общий документ', type='primary')
    identity = hashlib.sha256(upload.getvalue() + edited.to_csv(index=False).encode()).hexdigest()
    if st.session_state.get('talabot_batch_identity') != identity:
        st.session_state.pop('talabot_batch_result', None)
    if submitted:
        st.session_state.pop('talabot_batch_result', None)
        bar = st.progress(0.0)
        try:
            st.session_state.talabot_batch_result = generate_batch(df, edited, lambda done, total: bar.progress(done / total))
            st.session_state.talabot_batch_identity = identity
            st.success(f'Готов один Word для {len(edited)} групп.')
        except Exception as error:
            st.error(f'Не удалось создать документы: {error}')
    if st.session_state.get('talabot_batch_result'):
        st.download_button('Скачать общий Word', st.session_state.talabot_batch_result, file_name='Талабот_все_группы.docx', mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document', on_click='ignore')
