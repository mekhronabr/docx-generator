"""Rebuild the one-page Talabot form to match the supplied photograph.

Run with the bundled document Python runtime. The document uses in-flow
nested tables, so Word and PDF do not depend on floating-table positioning.
"""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parents[1]
INTRO = ('Куратори гуруҳ бо талаботи мазкур, ки дар саҳифаи аввал оварда шуда аст, '
         'моро шинос намуда моҳият ва мақсади онро шарҳ доданд. Мо донишҷӯёни гурӯҳ '
         'ба моҳияти талаботи мазкур сарфаҳм рафтем ва ӯҳдадор мешавем, ки онро риоя намоем.')


def paragraph(p, size=10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.alignment = align
    f = p.paragraph_format
    f.space_before = Pt(0)
    f.space_after = Pt(0)
    f.line_spacing = 1
    f.keep_with_next = False
    f.widow_control = False
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold, run.italic = bold, italic
    return p


def slot(p, tag, visible='', size=10, bold=False, italic=False, underline=False):
    run = p.add_run(visible)
    run.font.name, run.font.size = 'Times New Roman', Pt(size)
    run.bold, run.italic, run.underline = bold, italic, underline
    sdt = OxmlElement('w:sdt')
    props = OxmlElement('w:sdtPr')
    field = OxmlElement('w:tag')
    field.set(qn('w:val'), tag)
    props.append(field)
    sdt.append(props)
    content = OxmlElement('w:sdtContent')
    p._p.replace(run._r, sdt)
    content.append(run._r)
    sdt.append(content)


def table_style(table, widths, bordered=False):
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Mm(width)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for tag in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + tag)
        edge.set(qn('w:val'), 'single' if bordered else 'nil')
        edge.set(qn('w:sz'), '4')
        edge.set(qn('w:color'), '777777')
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for tag, value in [('top', 0), ('bottom', 0), ('left', 45 if bordered else 0), ('right', 45 if bordered else 0)]:
        edge = OxmlElement('w:' + tag)
        edge.set(qn('w:w'), str(value))
        edge.set(qn('w:type'), 'dxa')
        margins.append(edge)
    props.append(margins)
    for row in table.rows:
        cant = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant)
        for cell, width in zip(row.cells, widths):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph(cell.paragraphs[0])


def spacer(doc, size=3):
    p = doc.add_paragraph()
    paragraph(p, size=1)
    p.paragraph_format.line_spacing = Pt(size)
    p.paragraph_format.space_after = Pt(0)


def paired(doc, absent=False):
    outer = doc.add_table(rows=1, cols=3)
    table_style(outer, (92.5, 5, 92.5))
    for side, outer_cell in enumerate((outer.cell(0, 0), outer.cell(0, 2))):
        # Remove the initial empty paragraph; leave only the required tiny tail.
        initial = outer_cell.paragraphs[0]._p
        outer_cell._tc.remove(initial)
        count = 5 if absent else 20
        widths = (9, 83.5) if absent else (9, 60, 23.5)
        table = outer_cell.add_table(rows=count + 1, cols=len(widths))
        table_style(table, widths, True)
        headers = ('№', 'Ному насаб') if absent else ('№', 'Ному насаб', 'Имзо')
        for cell, label in zip(table.rows[0].cells, headers):
            cell.text = label
            paragraph(cell.paragraphs[0], 11, True, align=WD_ALIGN_PARAGRAPH.CENTER)
        table.rows[0].height, table.rows[0].height_rule = Pt(14), WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, row in enumerate(table.rows[1:], 1):
            number = side * count + index
            row.height, row.height_rule = Pt(25 if not absent else 14), WD_ROW_HEIGHT_RULE.AT_LEAST
            row.cells[0].text = f'{number}.'
            paragraph(row.cells[0].paragraphs[0], 10)
            p = paragraph(row.cells[1].paragraphs[0], 10)
            slot(p, f'{"absent" if absent else "student"}_{number}')
        for p in outer_cell.paragraphs:
            paragraph(p, 1)
            p.paragraph_format.line_spacing = Pt(1)


def build():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Mm(10)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    top = doc.add_table(rows=1, cols=2)
    table_style(top, (25, 165))
    p = paragraph(top.cell(0, 0).paragraphs[0], align=WD_ALIGN_PARAGRAPH.CENTER)
    slot(p, 'form_code', 'ФЭ_11.', size=12, bold=True)
    top.cell(0, 1).text = INTRO
    paragraph(top.cell(0, 1).paragraphs[0], 12, True, True, WD_ALIGN_PARAGRAPH.JUSTIFY)
    spacer(doc, 4)
    paired(doc)
    spacer(doc, 6)
    p = paragraph(doc.add_paragraph(), 11, italic=True)
    p.add_run('Ба ҳайъати асосии гурӯҳи академии  ')
    slot(p, 'group', '_______________', size=11, bold=True, italic=True, underline=True)
    p.add_run('  моҳияти талаботи мазкур санаи ')
    slot(p, 'date', '__________', size=11, italic=True, underline=True)
    p.add_run(' дар доираи дарси кураторӣ фаҳмонида шуд. Дар дарс донишҷӯёни зерин иштирок надоштанд.')
    for run in p.runs:
        run.font.name, run.font.size, run.italic = 'Times New Roman', Pt(11), True
    spacer(doc, 4)
    paired(doc, absent=True)
    spacer(doc, 8)
    signatures = doc.add_table(rows=2, cols=3)
    table_style(signatures, (70, 40, 80))
    for row, label, tag in zip(signatures.rows, ('Куратори гурӯҳи академӣ:', 'Муовини декан оид ба тарбия:'), ('curator', 'deputy')):
        row.height, row.height_rule = Pt(27), WD_ROW_HEIGHT_RULE.AT_LEAST
        row.cells[0].text = label
        paragraph(row.cells[0].paragraphs[0], 11, True, True)
        row.cells[1].text = '___________________'
        paragraph(row.cells[1].paragraphs[0], 11, align=WD_ALIGN_PARAGRAPH.CENTER)
        p = paragraph(row.cells[2].paragraphs[0], 11, align=WD_ALIGN_PARAGRAPH.CENTER)
        slot(p, tag, '_____________________________', size=11, bold=True, italic=True, underline=True)
        border = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:color'), '777777')
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)
    doc.save(BASE / 'template_talabot.docx')


if __name__ == '__main__':
    build()
